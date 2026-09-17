"""Idempotence + definition-integrity: conforming an already-conformed document must introduce no
further unexplained changes, and no pre-existing numbering definition may change meaning underneath the
references that depend on it."""
import os
import sys
import tempfile
import zipfile
import re

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'synthetic'))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))

import lib                                            # noqa: E402
from docx import Document                             # noqa: E402
from conformer.engine import Conformer                # noqa: E402

TEMPLATE = os.path.join(os.path.dirname(__file__), '..', 'src', 'conformer', 'assets', 'template.dotx')


def _sample():
    d = Document(lib.li_base_docx())
    for p in list(d.paragraphs):
        p._element.getparent().remove(p._element)
    d.add_paragraph('BACKGROUND', style='Heading1')
    d.add_paragraph('Body paragraph establishing a fact.', style='BodyText')
    for i in range(3):
        d.add_paragraph(f'Numbered clause {i}.', style='NumberedParagraph')
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).paragraphs[0].add_run('Header')
    t.cell(1, 0).paragraphs[0].add_run('Data')
    p = os.path.join(tempfile.mkdtemp(), 'sample.docx')
    d.save(p)
    return p


def _body(docx_path):
    doc = zipfile.ZipFile(docx_path).read('word/document.xml').decode('utf8')
    return re.search(r'<w:body>.*</w:body>', doc, re.S).group(0)


def _conform_to(path, out):
    c = Conformer(TEMPLATE, path)
    c.run()
    c.save(out)
    return c


def test_second_conform_is_idempotent():
    src = _sample()
    d = tempfile.mkdtemp()
    out1 = os.path.join(d, 'out1.docx')
    out2 = os.path.join(d, 'out2.docx')
    _conform_to(src, out1)
    c2 = _conform_to(out1, out2)
    # the twice-conformed body is identical to the once-conformed body
    assert _body(out1) == _body(out2)
    # and the second run reports no numbering flips or definition-integrity violations
    assert c2.numbering_report() == []
    assert c2.definition_integrity_report() == []


def test_definition_integrity_holds_on_conform():
    src = _sample()
    out = os.path.join(tempfile.mkdtemp(), 'out.docx')
    c = _conform_to(src, out)
    # no pre-existing numbering definition changed meaning underneath its references
    assert c.definition_integrity_report() == [], c.definition_integrity_report()


def test_outcome_report_three_sections():
    src = _sample()
    c = Conformer(TEMPLATE, src)
    c.run()
    rep = c.outcome_report()
    assert set(rep) == {'preservation', 'conformance', 'unresolved'}
    assert rep['conformance']['numbering_flips'] == []
    assert rep['conformance']['tables_failing_effective_format'] == []
