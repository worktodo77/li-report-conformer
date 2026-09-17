"""Numbering-repair verification: conforming must not silently flip a list's MEANING (number<->bullet),
and the resolution-based numbering_report() must catch it when it does. Complements test_numbering_graph
(resolver) and test_numbering_preserve (_keep_numpr containment)."""
import os
import sys
import tempfile

SYN = os.path.join(os.path.dirname(__file__), '..', 'synthetic')
sys.path.insert(0, SYN)
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))

import lib                                            # noqa: E402
from docx import Document                             # noqa: E402
from conformer.engine import Conformer                # noqa: E402

TEMPLATE = os.path.join(os.path.dirname(__file__), '..', 'src', 'conformer', 'assets', 'template.dotx')
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _numbering(fmt):
    return (f'<w:numbering {W}>'
            f'<w:abstractNum w:abstractNumId="40"><w:lvl w:ilvl="0"><w:start w:val="1"/>'
            f'<w:numFmt w:val="{fmt}"/><w:lvlText w:val="%1."/></w:lvl></w:abstractNum>'
            f'<w:num w:numId="7"><w:abstractNumId w:val="40"/></w:num></w:numbering>')


def _styles(fmt_holder='<w:styles ' + W + '></w:styles>'):
    return fmt_holder


def test_numbering_report_flags_meaning_flip():
    # A style resolves to a decimal list originally, a bullet list after — the Warhoe failure. The
    # resolution-based report must flag it (the preservation gate cannot, being numbering-blind).
    styles = (f'<w:styles {W}><w:style w:type="paragraph" w:styleId="NumberedParagraph">'
              f'<w:name w:val="Numbered Paragraph"/>'
              f'<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="7"/></w:numPr></w:pPr></w:style></w:styles>')
    f = Conformer.__new__(Conformer)
    f._orig_num0 = _numbering('decimal'); f._orig_styles0 = styles
    f.num = _numbering('bullet'); f.styles = styles
    rep = Conformer.numbering_report(f)
    assert len(rep) == 1
    assert rep[0]['style'] == 'NumberedParagraph'
    assert rep[0]['before'] == 'decimal' and rep[0]['after'] == 'bullet'
    assert rep[0]['meaning_flip'] is True


def test_numbering_report_empty_when_unchanged():
    styles = f'<w:styles {W}><w:style w:type="paragraph" w:styleId="X"><w:name w:val="X"/>' \
             f'<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="7"/></w:numPr></w:pPr></w:style></w:styles>'
    f = Conformer.__new__(Conformer)
    f._orig_num0 = _numbering('decimal'); f._orig_styles0 = styles
    f.num = _numbering('decimal'); f.styles = styles
    assert Conformer.numbering_report(f) == []


def test_conforming_does_not_flip_a_numbered_list():
    # Integration: a real conform of a document with a style-based numbered list must leave its meaning
    # intact (0 flips) and keep it resolving to a number, not a bullet.
    d = Document(lib.li_base_docx())
    for p in list(d.paragraphs):
        p._element.getparent().remove(p._element)
    d.add_paragraph('BACKGROUND', style='Heading1')
    for i in range(3):
        d.add_paragraph(f'Numbered clause {i} establishing a fact for the record.', style='NumberedParagraph')
    path = os.path.join(tempfile.mkdtemp(), 'numlist.docx')
    d.save(path)

    c = Conformer(TEMPLATE, path)
    c.run()
    assert c.numbering_report() == [], c.numbering_report()
    from conformer.numbering import NumberingGraph
    g = NumberingGraph(c.num, c.styles)
    pn = g.style_numpr('NumberedParagraph')
    assert pn is not None
    assert g.is_bullet(*pn) is False        # still a number, not a bullet
    ok, _ = c.validate_output()
    assert ok
