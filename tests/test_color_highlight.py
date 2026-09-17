"""Colour + highlight judgment calls: colour matching the paragraph style is stripped silently; colour
deviating from the style, and every highlight, become accept/skip judgment calls (accept normalises the
colour / removes the highlight; skip leaves it)."""
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'synthetic'))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))

import lib                                            # noqa: E402
from docx import Document                             # noqa: E402
from conformer.engine import Conformer                # noqa: E402

TEMPLATE = os.path.join(os.path.dirname(__file__), '..', 'src', 'conformer', 'assets', 'template.dotx')


def _doc_with(runs):
    d = Document(lib.li_base_docx())
    for p in list(d.paragraphs):
        p._element.getparent().remove(p._element)
    d.add_paragraph('BACKGROUND', style='Heading1')
    for text, rpr in runs:
        p = d.add_paragraph('', style='BodyText')
        r = p.add_run(text)
        if rpr:
            lib.add_direct_rpr(r, rpr)
    path = os.path.join(tempfile.mkdtemp(), 'ch.docx')
    d.save(path)
    return path


def _body(fresh):
    return fresh.head + ''.join(fresh.items) + fresh.tail


def test_deviating_colour_becomes_a_judgment_call():
    path = _doc_with([('This body text is blue.', '<w:color w:val="0563C1"/>')])
    c = Conformer(TEMPLATE, path)
    calls = c.analyze()
    color = [j for j in calls if j.kind == 'color']
    assert len(color) == 1, [j.kind for j in calls]
    # accept -> colour normalised away
    accept = c.apply_with_decisions({j.id: 'accept' for j in calls})
    assert '0563C1' not in _body(accept)
    # skip -> colour kept
    skip = c.apply_with_decisions({j.id: 'skip' for j in calls})
    assert '0563C1' in _body(skip)


def test_highlight_becomes_a_judgment_call():
    path = _doc_with([('Placeholder name here.', '<w:highlight w:val="yellow"/>')])
    c = Conformer(TEMPLATE, path)
    calls = c.analyze()
    hi = [j for j in calls if j.kind == 'highlight']
    assert len(hi) == 1
    accept = c.apply_with_decisions({j.id: 'accept' for j in calls})
    assert '<w:highlight' not in _body(accept)
    skip = c.apply_with_decisions({j.id: 'skip' for j in calls})
    assert '<w:highlight w:val="yellow"' in _body(skip)


def test_colour_matching_style_is_stripped_silently_no_call():
    # a run coloured to MATCH the BodyText style's own colour is redundant: removed with no judgment call
    probe = Conformer(TEMPLATE, _doc_with([('x', None)]))
    style_col = probe._effective_style_color('BodyText', probe._style_color_map()) or '000000'
    path = _doc_with([(f'Matches the style ({style_col}).', f'<w:color w:val="{style_col}"/>')])
    c = Conformer(TEMPLATE, path)
    calls = c.analyze()
    assert [j for j in calls if j.kind == 'color'] == []      # no deviation -> no call
    fresh = c.apply_with_decisions({})
    para = [it for it in fresh.items if 'Matches the style' in it]
    assert para and '<w:color' not in para[0]                 # redundant direct colour removed from it
