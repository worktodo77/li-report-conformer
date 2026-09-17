"""Property-specific direct-formatting policy: MEANING-BEARING formatting (super/subscript such as the
H2S subscript, strike, hidden) and REVIEW formatting (highlight) survive, while house-controlled
appearance (font/size/spacing/underline/caps) is normalised. Covers the shared filter used by both body
text (strip_direct) and footnotes (fix_footnotes)."""
import os
import sys
import tempfile

SYN = os.path.join(os.path.dirname(__file__), '..', 'synthetic')
sys.path.insert(0, SYN)
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))

from conformer.engine import keep_rpr_children       # noqa: E402


def _tags(rpr_inner, **kw):
    import re
    return [re.match(r'<w:(\w+)', c).group(1) for c in keep_rpr_children(rpr_inner, **kw)]


def test_subscript_survives_house_font_and_size_stripped():
    # H2S: the "2" run carries a subscript plus a body font + size. Subscript is meaning; font/size are
    # house appearance.
    rpr = '<w:rFonts w:ascii="Arial" w:hAnsi="Arial"/><w:sz w:val="20"/><w:vertAlign w:val="subscript"/>'
    assert _tags(rpr) == ['vertAlign']


def test_superscript_survives():
    assert 'vertAlign' in _tags('<w:vertAlign w:val="superscript"/><w:sz w:val="18"/>')


def test_strike_and_hidden_survive():
    assert _tags('<w:strike/><w:sz w:val="22"/>') == ['strike']
    # colour is retained by the filter (the colour/highlight pass decides it); size is house → gone
    assert _tags('<w:vanish/><w:specVanish/><w:sz w:val="22"/>') == ['vanish', 'specVanish']
    assert set(_tags('<w:vanish/><w:color w:val="FF0000"/>')) == {'vanish', 'color'}


def test_highlight_survives():
    assert _tags('<w:highlight w:val="yellow"/><w:spacing w:val="2"/>') == ['highlight']


def test_symbol_font_survives_body_font_stripped():
    assert _tags('<w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings"/>') == ['rFonts']   # glyph meaning
    assert _tags('<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>') == []                # body font gone


def test_house_appearance_is_stripped():
    rpr = '<w:sz w:val="24"/><w:szCs w:val="24"/><w:u w:val="single"/><w:caps/><w:spacing w:val="4"/><w:position w:val="6"/>'
    assert _tags(rpr) == []


def test_emphasis_and_charstyle_kept():
    assert set(_tags('<w:rStyle w:val="Emphasis"/><w:b/><w:i/>')) == {'rStyle', 'b', 'i'}


def test_footnote_reference_italic_dropped():
    assert _tags('<w:i/><w:iCs/>', is_fnref=True) == []
    assert set(_tags('<w:i/><w:iCs/>', is_fnref=False)) == {'i', 'iCs'}


def test_colour_retained_by_filter_for_the_colour_pass():
    # colour is no longer stripped by the run filter — it is preserved here and adjudicated later by the
    # colour/highlight pass (redundant-vs-style stripped silently; deviations become judgment calls).
    assert _tags('<w:color w:val="FFFFFF"/><w:b/>') == ['color', 'b']
    assert _tags('<w:color w:val="0563C1"/><w:sz w:val="20"/>') == ['color']


def test_h2s_subscript_survives_full_conform():
    # End-to-end: a paragraph "H2S" with the 2 subscripted must keep its subscript after conforming.
    import lib
    from docx import Document
    from conformer.engine import Conformer
    TEMPLATE = os.path.join(os.path.dirname(__file__), '..', 'src', 'conformer', 'assets', 'template.dotx')
    d = Document(lib.li_base_docx())
    for p in list(d.paragraphs):
        p._element.getparent().remove(p._element)
    d.add_paragraph('BACKGROUND', style='Heading1')
    p = d.add_paragraph('The gas ', style='BodyText')
    p.add_run('H')
    r2 = p.add_run('2')
    lib.add_direct_rpr(r2, '<w:vertAlign w:val="subscript"/><w:rFonts w:ascii="Arial"/><w:sz w:val="18"/>')
    p.add_run('S was detected.')
    path = os.path.join(tempfile.mkdtemp(), 'h2s.docx')
    d.save(path)

    c = Conformer(TEMPLATE, path)
    c.run()
    out = c.head + ''.join(c.items) + c.tail
    assert '<w:vertAlign w:val="subscript"/>' in out    # meaning kept
    assert 'w:ascii="Arial"' not in out or 'Wingdings' in out  # body font normalised away
