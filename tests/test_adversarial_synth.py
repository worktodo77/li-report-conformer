"""Adversarial synthetic constructs (Gate-A companion suite).

Small, targeted fixtures for the hard tracked-change / structural constructs that should NOT be
forced into the long reports: paragraph-mark deletion, a move, a table-cell revision, a tracked
table row, a corrupted style definition (Claire's real complaint), and a comment anchored wholly
inside a would-be-removed region. Each asserts the engine's expected disposition — preserve, hold,
or surface-as-restriction — with preservation verified clean. Fictitious content only.
"""
import os
import sys
import tempfile

import pytest

SYN = os.path.join(os.path.dirname(__file__), '..', 'synthetic')
sys.path.insert(0, SYN)
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))

import lib                                          # noqa: E402
from docx import Document                           # noqa: E402
from docx.oxml.ns import qn, nsdecls                # noqa: E402
from docx.oxml import parse_xml                     # noqa: E402
from conformer.engine import Conformer              # noqa: E402

TEMPLATE = os.path.join(os.path.dirname(__file__), '..', 'src', 'conformer', 'assets', 'template.dotx')


def _doc():
    d = Document(lib.li_base_docx())
    for p in list(d.paragraphs):
        p._element.getparent().remove(p._element)
    d.add_paragraph('BACKGROUND', style='Heading1')
    return d


def _save(d):
    p = os.path.join(tempfile.mkdtemp(), 'adv.docx')
    d.save(p)
    return p


def _conform(path):
    c = Conformer(TEMPLATE, path)
    c.run()
    clean, disc = c.verify_preservation()
    ok, _ = c.validate_output()
    return c, clean, disc, ok


def test_paragraph_mark_deletion_preserved():
    d = _doc()
    ids = lib.Ids()
    p = d.add_paragraph('This paragraph mark is deleted so it merges with the next.',
                        style='NumberedParagraph')
    d.add_paragraph('Following paragraph.', style='NumberedParagraph')
    lib.mark_paragraph_deleted(p, ids, 'Ann', '2026-01-01T00:00:00Z')
    c, clean, disc, ok = _conform(_save(d))
    assert clean, disc
    assert ok
    assert 'w:author="Ann"' in ''.join(c.items)


def test_paragraph_mark_only_deletion_preserved():
    d = _doc(); ids = lib.Ids()
    p = d.add_paragraph('Visible text remains while only the boundary is deleted.', style='NumberedParagraph')
    rec = lib.mark_paragraph_mark_deleted(p, ids, 'Ann', '2026-01-06T00:00:00Z')
    assert rec['kind'] == 'para_mark_del'
    c, clean, disc, ok = _conform(_save(d))
    assert clean, disc
    assert ok
    out = ''.join(c.items)
    assert 'w:del' in out and 'Visible text remains' in out

def test_move_from_to_preserved():
    d = _doc()
    ids = lib.Ids()
    a = d.add_paragraph(style='NumberedParagraph'); ra = a.add_run('moved text here')
    b = d.add_paragraph(style='NumberedParagraph'); rb = b.add_run('moved text here')
    lib.move_runs(ra, rb, ids, 'Bob', '2026-01-02T00:00:00Z', 'move1')
    c, clean, disc, ok = _conform(_save(d))
    assert clean, disc
    assert ok
    out = ''.join(c.items)
    assert '<w:moveFrom' in out and '<w:moveTo' in out


def test_table_cell_revision_surfaced_and_preserved():
    d = _doc()
    ids = lib.Ids()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).paragraphs[0].add_run('Header')
    t.cell(1, 0).paragraphs[0].add_run('Data')
    lib.add_cell_revision(t.cell(1, 1), 'del', ids, 'Cy', '2026-01-03T00:00:00Z')
    path = _save(d)
    c = Conformer(TEMPLATE, path)
    # the cellDel construct is surfaced as an explicit restriction, not silently dropped
    assert any(u[1] == 'cellDel' for u in c.revision_ledger.unsupported)
    c.run()
    clean, disc = c.verify_preservation()
    assert clean, disc


def test_tracked_table_row_preserved():
    d = _doc()
    ids = lib.Ids()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).paragraphs[0].add_run('H')
    t.cell(1, 0).paragraphs[0].add_run('R')
    # w:ins inside trPr marks the row itself inserted
    tr = t.rows[1]._tr
    trPr = tr.find(qn('w:trPr'))
    if trPr is None:
        trPr = parse_xml('<w:trPr %s/>' % nsdecls('w')); tr.insert(0, trPr)
    trPr.append(parse_xml('<w:ins %s w:id="7001" w:author="Dana" w:date="2026-01-04T00:00:00Z"/>'
                          % nsdecls('w')))
    c, clean, disc, ok = _conform(_save(d))
    assert clean, disc
    assert ok
    assert 'w:author="Dana"' in ''.join(c.items)


def test_corrupted_list_bullet_style_repaired():
    d = _doc()
    assert lib.corrupt_style_def(d, 'ListBullet')          # hollow the LI List Bullet definition
    d.add_paragraph('A bullet item that relies on the List Bullet style.', style='ListBullet')
    path = _save(d)
    c = Conformer(TEMPLATE, path)
    c.run()
    # _repair_styles overwrites the corrupt definition from the template (restores real formatting)
    import re
    m = re.search(r'<w:style\b[^>]*w:styleId="ListBullet".*?</w:style>', c.styles, re.S)
    assert m and ('<w:numPr' in m.group(0) or '<w:ind ' in m.group(0)), 'List Bullet not repaired'
    ok, _ = c.validate_output()
    assert ok


def test_corrupted_table_style_repaired():
    # A document with a table but missing the house table style ends up with a correct Grid Table 4 /
    # "LI Table": present in the output styles, with the black grid and the CONCRETE teal header (no navy,
    # no theme-dependent fill), so it renders teal regardless of the document's accent5.
    d = _doc()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).paragraphs[0].add_run('x')
    path = _save(d)
    c = Conformer(TEMPLATE, path)
    c.run()
    import re
    m = re.search(r'<w:style\b[^>]*w:styleId="GridTable4".*?</w:style>', c.styles, re.S)
    assert m and '<w:tblBorders' in m.group(0), 'GridTable4 ("LI Table") not present/repaired from template'
    assert 'B6DDE8' in m.group(0) and '054F8A' not in m.group(0) and 'themeFill' not in m.group(0)


def test_comment_inside_removed_region_preserved():
    d = _doc()
    ids = lib.Ids()
    p = d.add_paragraph(style='NumberedParagraph')
    r = p.add_run('commented text that is also inside a deletion')
    lib.add_comment(d, r, 'Check this deleted passage.', 'Ed', 'ED')
    lib.wrap_run_del(r, ids, 'Ed', '2026-01-05T00:00:00Z')       # the commented run is deleted
    c, clean, disc, ok = _conform(_save(d))
    assert clean, disc                                            # comment anchor + body preserved
    assert ok
