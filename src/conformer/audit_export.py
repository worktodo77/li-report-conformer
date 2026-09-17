"""Full audit log of every change the conformer made, exportable as CSV, Excel or Word.

`build_records()` turns a finished (applied) Conformer into a summary dict + a flat list of per-change
records following the agreed schema; `write_csv` / `write_xlsx` / `write_docx` render those. Everything
is derived from what actually ran, so the log reflects the real output, not intentions."""
import os, csv, hashlib, datetime

# category key -> (change type, element, basis/rule) for the skippable/structural conformance passes.
_CAT_META = {
    'styles-repair': ('Repair corrupt style definition', 'Style definition', 'Template style repair'),
    'tables':        ('Apply LI table style', 'Table', 'LI table style'),
    'footnotes':     ('Normalize footnote', 'Footnote', 'Footnote formatting'),
    'empty-paras':   ('Delete empty paragraph', 'Paragraph', 'Structural policy #2 (empty paragraphs)'),
    'page-breaks':   ('Remove manual page break', 'Paragraph', 'Structural policy #4 (manual page breaks)'),
    'pdf-merge':     ('Merge PDF line-break split', 'Paragraph text', 'Structural policy #1 (PDF line splits)'),
}

_COLUMNS = ['#', 'Category', 'Change type', 'Element', 'Location', 'Before', 'After',
            'Basis / rule', 'Disposition', 'Reason', 'Tracked-change safe']


def _sha256(path):
    try:
        h = hashlib.sha256()
        with open(path, 'rb') as fh:
            for chunk in iter(lambda: fh.read(1 << 16), b''):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return '(unavailable)'


def _structural_meta(msg):
    """Change type / element / basis for a conformance log line that has no explicit category."""
    m = msg.lower()
    if 'caption' in m:
        return ('Rebuild caption field', 'Caption', 'STYLEREF/SEQ caption fields')
    if 'cross-reference' in m or 'ref field' in m:
        return ('Rebuild cross-reference', 'Cross-reference', 'REF fields')
    if 'floating picture' in m or 'inline figure' in m or 'picture' in m:
        return ('Convert floating image to inline', 'Image', 'Figure layout')
    if 'section break' in m or 'landscape' in m or 'portrait' in m:
        return ('Restore section break', 'Section', 'Section properties')
    if 'wrapper table' in m or 'unwrapped' in m:
        return ('Unwrap wrapper table', 'Table', 'Structural policy #3')
    if 'column' in m:
        return ('Remove empty column(s)', 'Table', 'Structural policy #7')
    if 'heading' in m:
        return ('Split merged heading', 'Paragraph', 'Structural policy #6')
    if 'figure' in m:
        return ('Figure conformance', 'Image', 'Figure layout')
    if 'level' in m:
        return ('Correct list level', 'Paragraph', 'List levels')
    if 'style' in m or 'numbering' in m:
        return ('Replace styles/numbering', 'Style definition', 'Template styles')
    return ('Formatting fix', 'Paragraph', 'LI template')


def _section_index(fresh):
    """List of (body_index, heading text) for the nearest-heading location lookups."""
    heads = []
    try:
        for i in range(fresh.n()):
            st = fresh.style(i)
            if st in ('Heading1', 'Heading2'):
                t = (fresh.text(i) or '').strip()
                if t:
                    heads.append((i, t[:60]))
    except Exception:
        pass
    return heads


def _locate(heads, item):
    if item is None or item < 0:
        return 'Whole document'
    sec = ''
    for idx, title in heads:
        if idx <= item:
            sec = title
        else:
            break
    return (f'{sec} · ¶{item + 1}' if sec else f'¶{item + 1}')


def build_records(fresh, source_path, output_path, decisions_log=None,
                  source_sha=None, output_sha=None):
    preserve = getattr(fresh, 'disposition', None) == 'preserve'
    heads = _section_index(fresh)
    tracked_safe = 'Yes (gate-verified)' if preserve else 'n/a (no tracked changes)'
    records = []

    def add(category, ctype, element, item, before, after, basis, disp='Applied', reason='', safe=None):
        records.append({
            'Category': category, 'Change type': ctype, 'Element': element,
            'Location': _locate(heads, item), 'Before': before or '', 'After': after or '',
            'Basis / rule': basis, 'Disposition': disp, 'Reason': reason,
            'Tracked-change safe': tracked_safe if safe is None else safe,
        })

    # 1) Conformance + structural fixes (engine log)
    for e in getattr(fresh, 'log', []):
        msg = e.get('msg')
        if not msg:
            continue
        cat = e.get('cat')
        if cat in _CAT_META:
            ctype, element, basis = _CAT_META[cat]
            category = 'Conformance fix'
        else:
            low = msg.lower()
            if low.startswith('typography normalised') or low.startswith('li house style applied'):
                continue   # cosmetic summary — the per-edit detail is added below
            ctype, element, basis = _structural_meta(msg)
            category = 'Conformance fix'
        add(category, ctype, element, e.get('item', -1), '', msg, basis)

    # 2) Text cleanup edits (typography + house style), incl. any the user skipped
    for ed in getattr(fresh, 'pending_edits', []):
        if ed['before'] == ed['after']:
            continue
        kind = ed.get('kind')
        ctype = 'Typography' if kind == 'typo' else ('House style' if kind == 'house' else ed.get('label', 'Text edit'))
        basis = 'LI house style' if kind == 'house' else 'LI typography'
        skipped = ed.get('skipped')
        add('Text cleanup', ctype, 'Paragraph text', ed.get('item', -1), ed['before'], ed['after'],
            basis, disp='Skipped by user' if skipped else 'Applied',
            reason='User skipped' if skipped else '')

    # 3) Judgment-call decisions (from the review UI)
    for d in (decisions_log or []):
        add('Judgment call', 'Structural (interactive)', 'Paragraph', -1,
            d.get('text', ''), d.get('action', ''), 'Structural-change policy (ASK)',
            disp=d.get('status', '').title() or 'Applied')

    # 4) Conformance categories the user skipped (didn't run, so no log line)
    for key, val in (getattr(fresh, 'decisions', None) or {}).items():
        if isinstance(key, str) and key.startswith('conf:') and val == 'skip':
            cat = key.split(':', 1)[1]
            ctype, element, basis = _CAT_META.get(cat, (cat, 'Paragraph', 'Conformance'))
            add('Conformance fix', ctype, element, -1, '', '(not applied)', basis,
                disp='Skipped by user', reason='User skipped this fix type')

    # 5) Figure-integrity advisories (don't change the file)
    for a in getattr(fresh, 'audit', []) or []:
        if isinstance(a, (tuple, list)) and len(a) == 2:
            lvl, m = a
        else:
            lvl, m = '', str(a)
        add('Figure integrity', f'Advisory [{lvl}]', 'Figure/Table', -1, '', m,
            'Figure numbering / Table of Figures', disp='Advisory', reason='Review in source',
            safe='n/a (advisory)')

    # 6) Passes rolled back by the preservation gate (recorded, not applied)
    for name, reason in getattr(fresh, 'exceptions', []) or []:
        add('Preservation', 'Pass rolled back', 'Document', -1, '', name,
            'Preservation gate', disp='Rolled back', reason=reason, safe='Yes (rolled back)')

    for i, r in enumerate(records, 1):
        r['#'] = i

    # Summary
    led = getattr(fresh, 'revision_ledger', None)
    s = led.summary() if led else {'total': 0, 'comments': 0, 'authors': []}
    try:
        clean, _disc = fresh.verify_preservation()
        verdict = 'Clean — every tracked change and comment preserved' if clean else 'VIOLATIONS — see rolled-back passes'
    except Exception:
        verdict = 'Not evaluated'
    from collections import Counter
    by_cat = Counter(r['Category'] for r in records)
    summary = {
        'source_name': os.path.basename(source_path),
        'source_sha256': source_sha or _sha256(source_path),
        'output_name': os.path.basename(output_path),
        'output_sha256': output_sha or _sha256(output_path),
        'template': os.path.basename(getattr(fresh, 'template_path', '') or ''),
        'run_at': datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'mode': 'Review-preserving' if preserve else 'Clean conform',
        'review_copy_label': getattr(fresh, 'REVIEW_COPY_LABEL', '') if preserve else '',
        'tracked_changes': s.get('total', 0),
        'comments': s.get('comments', 0),
        'authors': len(s.get('authors', []) or []),
        'preservation_verdict': verdict,
        'rolled_back': len(getattr(fresh, 'exceptions', []) or []),
        'total_changes': len(records),
        'by_category': dict(by_cat),
    }
    return summary, records


def _summary_rows(summary):
    return [
        ('Source file', summary['source_name']),
        ('Source SHA-256', summary['source_sha256']),
        ('Output file', summary['output_name']),
        ('Output SHA-256', summary['output_sha256']),
        ('Template', summary['template']),
        ('Run at', summary['run_at']),
        ('Mode', summary['mode']),
        ('Output label', summary['review_copy_label']),
        ('Tracked changes', summary['tracked_changes']),
        ('Comments', summary['comments']),
        ('Authors', summary['authors']),
        ('Preservation', summary['preservation_verdict']),
        ('Passes rolled back', summary['rolled_back']),
        ('Total changes logged', summary['total_changes']),
        ('By category', ', '.join(f'{k}: {v}' for k, v in summary['by_category'].items())),
    ]


def write_csv(path, summary, records):
    with open(path, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.writer(fh)
        w.writerow(['LI Report Conformer — audit log'])
        for k, v in _summary_rows(summary):
            w.writerow([k, v])
        w.writerow([])
        w.writerow(_COLUMNS)
        for r in records:
            w.writerow([r.get(c, '') for c in _COLUMNS])


def write_xlsx(path, summary, records):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    navy = PatternFill('solid', fgColor='005088')
    head_font = Font(bold=True, color='FFFFFF')
    bold = Font(bold=True)

    ws = wb.active; ws.title = 'Summary'
    ws['A1'] = 'LI Report Conformer — audit log'; ws['A1'].font = Font(bold=True, size=14)
    row = 3
    for k, v in _summary_rows(summary):
        ws.cell(row, 1, k).font = bold
        ws.cell(row, 2, str(v))
        row += 1
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 90

    def sheet(title, rows):
        s = wb.create_sheet(title[:31])
        for c, name in enumerate(_COLUMNS, 1):
            cell = s.cell(1, c, name); cell.font = head_font; cell.fill = navy
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        for ri, r in enumerate(rows, 2):
            for c, name in enumerate(_COLUMNS, 1):
                s.cell(ri, c, str(r.get(name, ''))).alignment = Alignment(vertical='top', wrap_text=True)
        widths = {'#': 5, 'Category': 15, 'Change type': 24, 'Element': 15, 'Location': 26,
                  'Before': 44, 'After': 44, 'Basis / rule': 26, 'Disposition': 16, 'Reason': 22,
                  'Tracked-change safe': 18}
        for c, name in enumerate(_COLUMNS, 1):
            s.column_dimensions[s.cell(1, c).column_letter].width = widths.get(name, 18)
        s.freeze_panes = 'A2'
        return s

    sheet('All changes', records)
    for cat, title in (('Conformance fix', 'Conformance fixes'), ('Text cleanup', 'Text cleanup'),
                       ('Judgment call', 'Judgment calls'), ('Figure integrity', 'Figure integrity'),
                       ('Preservation', 'Preservation')):
        rows = [r for r in records if r['Category'] == cat]
        if rows:
            sheet(title, rows)
    wb.save(path)


def write_docx(path, summary, records):
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc = Document()
    h = doc.add_heading('LI Report Conformer — audit log', level=0)
    try:
        h.runs[0].font.color.rgb = RGBColor(0x00, 0x50, 0x88)
    except Exception:
        pass

    doc.add_heading('Summary', level=1)
    st = doc.add_table(rows=0, cols=2); st.style = 'Light List Accent 1'
    for k, v in _summary_rows(summary):
        c = st.add_row().cells
        c[0].text = k; c[1].text = str(v)
        for p in c[0].paragraphs:
            for r in p.runs:
                r.font.bold = True

    doc.add_heading(f'Changes ({len(records)})', level=1)
    cols = ['#', 'Category', 'Change type', 'Element', 'Location', 'Before', 'After',
            'Basis / rule', 'Disposition', 'Tracked-change safe']
    t = doc.add_table(rows=1, cols=len(cols)); t.style = 'Light Grid Accent 1'
    for c, name in enumerate(cols):
        cell = t.rows[0].cells[c]; cell.text = name
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True; r.font.size = Pt(8)
    for rec in records:
        cells = t.add_row().cells
        for c, name in enumerate(cols):
            cells[c].text = str(rec.get(name, ''))
            for p in cells[c].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
    doc.save(path)
